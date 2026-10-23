"""Shared local PDF/image/PPTX reconstruction. Native text first, residual OCR second."""
from pathlib import Path
from contextlib import contextmanager
import sys,os,json,hashlib,tempfile,subprocess,shutil,re,math,statistics,copy
sys.dont_write_bytecode=True

SKILL=Path(__file__).resolve().parent.parent
sys.path[:0]=[str(SKILL/'vendor'),str(SKILL/'vendor/ocr')]
import pymupdf as fitz
from PIL import Image,ImageFont

VERSION='2.0'
def norm(text):return re.sub(r'\s+','',text)
def digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def run_local(command,timeout=300):
    flags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0
    proc=subprocess.Popen([str(x) for x in command],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,creationflags=flags)
    try:raw,_=proc.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        if os.name=='nt':subprocess.run(['taskkill','/PID',str(proc.pid),'/T','/F'],capture_output=True,creationflags=flags)
        else:proc.kill()
        proc.communicate();raise RuntimeError('This conversion process timed out; unrelated Office processes were not stopped.')
    if proc.returncode:raise RuntimeError(raw.decode('utf-8',errors='replace')[-4000:])
    return raw.decode('utf-8',errors='replace')

def render_office(source,outdir,profile):
    source=Path(source).resolve();outdir=Path(outdir).resolve();outdir.mkdir(parents=True,exist_ok=True)
    profile=Path(profile).resolve();profile.mkdir(parents=True,exist_ok=True)
    soffice=os.environ.get('PDF_CONVERT_SOFFICE') or shutil.which('soffice.com') or shutil.which('soffice')
    if not soffice and Path(r'C:\Program Files\LibreOffice\program\soffice.com').exists():soffice=r'C:\Program Files\LibreOffice\program\soffice.com'
    if not soffice:raise RuntimeError('LibreOffice is needed to read PPTX or render output; set PDF_CONVERT_SOFFICE.')
    target=outdir/(source.stem+'.pdf')
    if target.exists():raise FileExistsError(f'Render target already exists: {target}')
    run_local([soffice,f'-env:UserInstallation={profile.as_uri()}','--headless','--norestore','--convert-to','pdf','--outdir',outdir,source])
    if not target.exists() or target.stat().st_size==0:raise RuntimeError('LibreOffice returned without a new PDF.')
    return target

@contextmanager
def normalise(source,work):
    source=Path(source).resolve();work=Path(work);ext=source.suffix.lower()
    if not source.exists():raise FileNotFoundError(source)
    if ext=='.pdf':doc=fitz.open(source)
    elif ext in ['.png','.jpg','.jpeg','.bmp','.tif','.tiff','.webp']:
        im=Image.open(source);doc=fitz.open()
        from PIL import ImageSequence
        for i,frame in enumerate(ImageSequence.Iterator(im)):
            frame=frame.convert('RGB');path=work/f'input-{i}.png';frame.save(path)
            scale=min(960/frame.width,960/frame.height)
            page=doc.new_page(width=frame.width*scale,height=frame.height*scale)
            page.insert_image(page.rect,filename=str(path))
    elif ext in ['.pptx','.ppt']:
        pdf=render_office(source,work/'normalised',work/'lo-input-profile');doc=fitz.open(pdf)
    else:raise ValueError('Input must be PDF, an image, or PPT/PPTX.')
    try:
        if len(doc)==0:raise ValueError('Input contains no pages.')
        yield doc
    finally:doc.close()

def native_lines(page):
    return [dict(line,origin='native') for block in page.get_text('dict')['blocks'] if block['type']==0
            for line in block['lines'] if any(s['text'].strip() for s in line['spans'])]

def textless_page(page):
    doc=fitz.open();doc.insert_pdf(page.parent,from_page=page.number,to_page=page.number);p=doc[0]
    for line in native_lines(p):
        for span in line['spans']:
            if span['text'].strip():p.add_redact_annot(fitz.Rect(span['bbox']),fill=None)
    p.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE,graphics=fitz.PDF_REDACT_LINE_ART_NONE,text=fitz.PDF_REDACT_TEXT_REMOVE)
    return doc

def make_rapid():
    from rapidocr import RapidOCR,OCRVersion,ModelType
    models=SKILL/'vendor/ocr_models'
    manifest=models/'models.json'
    if not manifest.exists():raise RuntimeError('Offline OCR models are missing. No models will be downloaded by conversion.')
    params={'Global.model_root_dir':str(models),'Global.log_level':'warning','Global.text_score':0.60,
            'EngineConfig.onnxruntime.intra_op_num_threads':4,'EngineConfig.onnxruntime.inter_op_num_threads':1,
            'Det.limit_side_len':1280,'Det.limit_type':'max'}
    for item in json.loads(manifest.read_text(encoding='utf-8')):
        path=models/item['file']
        if not path.exists() or digest(path)!=item['sha256']:raise RuntimeError('Offline OCR model integrity failure: '+item['file'])
        params[item['part']+'.model_path']=str(path)
        params[item['part']+'.ocr_version']=OCRVersion.PPOCRV6 if item['version']=='PP-OCRv6' else OCRVersion.PPOCRV4
        params[item['part']+'.model_type']=ModelType.SMALL if item['version']=='PP-OCRv6' else ModelType.MOBILE
    return RapidOCR(params=params)

def rapid_lines(engine,image):
    cache=getattr(engine,'_pdf_layout_cache',{});key=digest(image)
    if key in cache:return copy.deepcopy(cache[key])
    result=engine(str(image),use_det=True,use_cls=True,use_rec=True)
    if result.boxes is None:return []
    output=[];im=Image.open(image).convert('RGB')
    for box,text,score in zip(result.boxes,result.txts,result.scores):
        xs=[float(p[0]) for p in box];ys=[float(p[1]) for p in box]
        rect=[min(xs),min(ys),max(xs),max(ys)];item={'text':str(text),'score':float(score),'box':rect}
        if score<.90:
            import numpy as np
            votes=[]
            for pad in (5,15):
                x0,y0,x1,y1=rect
                crop=im.crop((max(0,int(x0)-pad),max(0,int(y0)-pad),min(im.width,int(x1)+pad),min(im.height,int(y1)+pad)))
                crop=crop.resize((crop.width*3,crop.height*3))
                retry=engine(np.array(crop),use_det=False,use_cls=False,use_rec=True)
                if retry.txts and retry.scores:votes.append((str(retry.txts[0]),float(retry.scores[0])))
            if len(votes)==2 and votes[0][0]==votes[1][0] and min(v[1] for v in votes)>=.95:
                item.update(text=votes[0][0],score=min(v[1] for v in votes),refined_from=item['text'])
        # A differently coloured diagram arrow can be attached to a label.
        if re.search('[←→↑↓↔↕]',item['text']):
            import numpy as np
            array=np.array(im);color,_=infer_style(array,rect,1,item['text'])
            ink=np.array([(color>>16)&255,(color>>8)&255,color&255])
            x0,y0,x1,y1=map(int,rect);tile=array[y0:y1,x0:x1]
            cols=np.where((np.max(np.abs(tile.astype(float)-ink),axis=2)<85).sum(axis=0)>max(1,tile.shape[0]*.06))[0]
            groups=[]
            for col in cols:
                if not groups or col-groups[-1][-1]>(y1-y0)*.5:groups.append([int(col)])
                else:groups[-1].append(int(col))
            if len(groups)>1:
                ranked=sorted(groups,key=lambda g:g[-1]-g[0],reverse=True)
                if ranked[1][-1]-ranked[1][0] < (ranked[0][-1]-ranked[0][0])*.35:cols=np.array(ranked[0])
            if len(cols) and max(cols[0],tile.shape[1]-1-cols[-1])>(y1-y0)*.65:
                left=max(0,x0+int(cols[0])-2);right=min(im.width,x0+int(cols[-1])+3)
                crop=im.crop((left,max(0,y0-4),right,min(im.height,y1+4)))
                rec=engine(np.array(crop),use_det=False,use_cls=False,use_rec=True)
                plain=re.sub('[←→↑↓↔↕]','',item['text'])
                if rec.txts and rec.scores and norm(str(rec.txts[0]))==norm(plain) and float(rec.scores[0])>=.95:
                    rect=[left,y0,right,y1];item.update(text=str(rec.txts[0]),score=float(rec.scores[0]),box=rect)
        # Detectors can merge adjacent diagram labels across a wide arrow/gap.
        if rect[2]-rect[0]>5*(rect[3]-rect[1]):
            import numpy as np
            array=np.array(im);color,_=infer_style(array,rect,1,item['text'])
            ink=np.array([(color>>16)&255,(color>>8)&255,color&255])
            x0,y0,x1,y1=map(int,rect);tile=array[y0:y1,x0:x1]
            active=np.where((np.max(np.abs(tile.astype(float)-ink),axis=2)<85).sum(axis=0)>max(1,tile.shape[0]*.06))[0]
            groups=[]
            for col in active:
                if not groups or col-groups[-1][-1]>(y1-y0)*1.3:groups.append([int(col)])
                else:groups[-1].append(int(col))
            if 2<=len(groups)<=4 and min(g[-1]-g[0] for g in groups)>(y1-y0)*.7:
                parts=[]
                for g in groups:
                    left=max(0,x0+g[0]-2);right=min(im.width,x0+g[-1]+3)
                    crop=im.crop((left,max(0,y0-4),right,min(im.height,y1+4)))
                    rec=engine(np.array(crop),use_det=False,use_cls=False,use_rec=True)
                    if rec.txts and rec.scores:parts.append({'text':str(rec.txts[0]),'score':float(rec.scores[0]),'box':[left,y0,right,y1]})
                if len(parts)==len(groups) and min(p['score'] for p in parts)>=.90 and norm(''.join(p['text'] for p in parts))==norm(re.sub('[←→↑↓↔↕]','',item['text'])):
                    output.extend(parts);continue
        output.append(item)
    if len(cache)<64:cache[key]=copy.deepcopy(output)
    engine._pdf_layout_cache=cache
    return output

def graphic_candidate(page,text,box):
    """A short OCR token covering one simple filled polygon is likely a graphic."""
    if len(norm(text))>2:return False
    r=fitz.Rect(box)
    for draw in page.get_drawings():
        items=draw.get('items',[])
        if draw.get('fill') is None or not 3<=len(items)<=12 or not all(x[0]=='l' for x in items):continue
        shape=draw['rect'];intersection=shape&r
        if shape.get_area()>200 and intersection.get_area()/max(1,r.get_area())>.65 and shape.contains(r.tl+(r.br-r.tl)*.5):return True
    return False

def windows_lines(image,work):
    req=work/'windows-input.json';out=work/'windows-output.json'
    req.write_text(json.dumps([{'id':1,'image':str(Path(image).resolve())}]),encoding='utf-8')
    shell=Path(os.environ.get('SystemRoot',r'C:\Windows'))/'System32/WindowsPowerShell/v1.0/powershell.exe'
    run_local([shell,'-NoProfile','-ExecutionPolicy','Bypass','-File',SKILL/'scripts/windows_ocr.ps1','-InputJson',req,'-OutputJson',out])
    data=json.loads(out.read_text(encoding='utf-8-sig'));items=[]
    for line in data[0]['lines']:
        words=line['words']
        if not words:continue
        box=[min(w['box'][0] for w in words),min(w['box'][1] for w in words),max(w['box'][2] for w in words),max(w['box'][3] for w in words)]
        text=re.sub(r'(?<=[\u3400-\u9fff])\s+(?=[\u3400-\u9fff])','',line['text'])
        items.append({'text':text,'score':None,'box':box})
    return items

def infer_style(array,box,scale,text):
    import numpy as np
    import cv2
    x0,y0,x1,y1=map(int,box);h,w=array.shape[:2]
    x0=max(0,x0);y0=max(0,y0);x1=min(w,x1+1);y1=min(h,y1+1)
    crop=array[y0:y1,x0:x1,:3]
    if crop.size==0:return 0,12
    pad=4;outer=array[max(0,y0-pad):min(h,y1+pad),max(0,x0-pad):min(w,x1+pad),:3]
    ring=np.concatenate([outer[:pad].reshape(-1,3),outer[-pad:].reshape(-1,3),outer[:,:pad].reshape(-1,3),outer[:,-pad:].reshape(-1,3)])
    bg=np.median(ring,axis=0)
    # Thin stroke contrast distinguishes ink from broad coloured highlights.
    gray=cv2.cvtColor(crop,cv2.COLOR_RGB2GRAY)
    radius=max(3,min(19,(y1-y0)//2))|1
    kernel=cv2.getStructuringElement(cv2.MORPH_RECT,(radius,radius))
    contrast=cv2.morphologyEx(gray,cv2.MORPH_BLACKHAT if float(np.median(gray))>140 else cv2.MORPH_TOPHAT,kernel)
    fg=crop[contrast>55]
    if not len(fg):fg=crop[np.max(np.abs(crop.astype(float)-bg),axis=2)>65]
    if len(fg):
        quant=(fg//24)*24;colors,counts=np.unique(quant,axis=0,return_counts=True)
        eligible=counts>=max(3,len(fg)*.015)
        candidates=np.where(eligible)[0]
        if not len(candidates):candidates=np.array([counts.argmax()])
        luma=colors[candidates]@np.array([.299,.587,.114])
        main=colors[candidates[luma.argmin() if float(np.median(gray))>140 else luma.argmax()]]
        selected=fg[np.max(np.abs(quant.astype(float)-main),axis=1)==0];rgb=np.median(selected,axis=0).astype(int)
    else:rgb=np.array([0,0,0])
    fontpath=Path(r'C:\Windows\Fonts\msyh.ttc') if re.search(r'[\u3400-\u9fff]',text) else Path(r'C:\Windows\Fonts\arial.ttf')
    try:
        font=ImageFont.truetype(str(fontpath),100);b=font.getbbox(text)
        size=min((x1-x0)/max(1,b[2]-b[0]),(y1-y0)/max(1,b[3]-b[1]))*100/scale
    except OSError:size=(y1-y0)/scale
    return int(rgb[0])*65536+int(rgb[1])*256+int(rgb[2]),max(4,round(size*2)/2)

def reconstruct(page,work,engine,ocr='auto',backend='rapid',visuals='preserve',ocr_all_regions=False):
    import numpy as np
    lines=native_lines(page);w,h=page.rect.width,page.rect.height
    residual=textless_page(page);scale=min(2.5,2400/max(w,h))
    image=work/f'page-{page.number+1:03}.png'
    residual[0].get_pixmap(matrix=fitz.Matrix(scale,scale),alpha=False).save(str(image));residual.close()
    im=Image.open(image).convert('RGB');array=np.array(im)
    added=[];ignored=[];review=[]
    if ocr!='off':
        recognized=rapid_lines(engine,image) if backend=='rapid' else windows_lines(image,work)
        for item in recognized:
            text=item['text'].strip();box=item['box'];x0,y0,x1,y1=box
            # Brand logos and running footers remain independent graphics.
            corner=((x0+x1)/2 < w*scale*.20 and (y0+y1)/2 < h*scale*.14) or ((x0+x1)/2>w*scale*.75 and (y0+y1)/2>h*scale*.91)
            if not ocr_all_regions and corner and x1-x0<w*scale*.25 and y1-y0<h*scale*.06:
                ignored.append({'text':text,'reason':'peripheral-brand-or-footer'});continue
            if not text:continue
            if text.strip() in ['一','-','—','→','←','↑','↓'] and x1-x0>2*(y1-y0):
                ignored.append({'text':text,'reason':'shape-like'});continue
            if graphic_candidate(page,text,[v/scale for v in box]):
                ignored.append({'text':text,'reason':'simple-vector-graphic','bbox':[v/scale for v in box]});continue
            color,size=infer_style(array,box,scale,text)
            bbox=[v/scale for v in box]
            span={'text':text,'font':'Microsoft YaHei' if re.search(r'[\u3400-\u9fff]',text) else 'Arial','size':size,'color':color,'flags':0,'bbox':bbox}
            line={'bbox':bbox,'spans':[span],'origin':'ocr','score':item['score']}
            added.append(line)
            if item['score'] is None or item['score']<.90:review.append({'text':text,'score':item['score'],'bbox':bbox})
        if added:
            import cv2
            mask=np.zeros(array.shape[:2],dtype='uint8')
            for line in added:
                box=[int(round(v*scale)) for v in line['bbox']];x0,y0,x1,y1=box
                # Limit the glyph mask to the detected area; do not paint whole rows white.
                x0=max(0,x0-3);y0=max(0,y0-3);x1=min(mask.shape[1],x1+4);y1=min(mask.shape[0],y1+4)
                tile=array[y0:y1,x0:x1].astype(float)
                if not tile.size:continue
                color=line['spans'][0]['color'];ink=np.array([(color>>16)&255,(color>>8)&255,color&255])
                glyph=(np.max(np.abs(tile-ink),axis=2)<130).astype('uint8')*255
                if line['spans'][0]['size']>18 and max(ink)>140 and not re.search('[←→↑↓↔↕]',line['spans'][0]['text']):
                    ring=np.concatenate([tile[:2].reshape(-1,3),tile[-2:].reshape(-1,3),tile[:,:2].reshape(-1,3),tile[:,-2:].reshape(-1,3)])
                    bg=np.median(ring,axis=0)
                    if np.max(np.std(ring,axis=0))<18:glyph=(np.max(np.abs(tile-bg),axis=2)>15).astype('uint8')*255
                glyph=cv2.dilate(glyph,np.ones((3,3),dtype='uint8'))
                mask[y0:y1,x0:x1]=np.maximum(mask[y0:y1,x0:x1],glyph)
            cleaned=cv2.inpaint(array,mask,2,cv2.INPAINT_TELEA)
            Image.fromarray(cleaned).save(image)
    if visuals=='none':
        # Deliberate text-only mode, never the default.
        Image.new('RGB',im.size,'white').save(image)
    return lines+added,image,{'page':page.number+1,'size':[w,h],'native_text':''.join(s['text'] for l in lines for s in l['spans']),
        'native_lines':len(lines),'ocr_lines':len(added),'ocr_text':''.join(s['text'] for l in added for s in l['spans']),
        'ocr_regions':[l['bbox'] for l in added],'ocr_items':[{'text':l['spans'][0]['text'],'score':l['score'],'bbox':l['bbox']} for l in added],
        'review':review,'static_peripheral_regions':ignored,'background':visuals!='none'}

def textbox_geometry(line,w,h):
    x0,y0,x1,y1=line['bbox'];size=max(s['size'] for s in line['spans']);height=max(y1-y0,size*1.02)
    center=(x0+x1)/2
    align='center' if abs(center-w/2)<12 and x1-x0<w*.9 else ('right' if x1>w-60 and x0>w*.5 else 'left')
    bw=min((x1-x0)*1.4+30,w)
    x=center-bw/2 if align=='center' else (x1+8-bw if align=='right' else x0)
    x=max(0,min(x,w-bw));bw=min(bw,w-x)
    return x,max(y0-2,0),max(bw,1),min(height*1.6+6,h-max(y0-2,0)),height,align

def convert(source,output,kind,ocr='auto',backend='rapid',visuals='preserve',ocr_all_regions=False):
    import time
    from docx import Document
    from docx.shared import Pt as WPt,RGBColor
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn,nsdecls
    from docx.enum.section import WD_SECTION
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.dml.color import RGBColor as PColor
    from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
    import pdf2docx_editable as word
    import pdf2docx_image as pictures
    import pdf2pptx as ppt
    source=Path(source).resolve();output=Path(output).resolve()
    if source==output:raise ValueError('Output must not overwrite the source file.')
    if output.suffix.lower()!='.'+kind:raise ValueError('Output extension must match '+kind)
    output.parent.mkdir(parents=True,exist_ok=True)
    t=time.perf_counter();report={'version':VERSION,'source':str(source),'source_sha256':digest(source),'kind':kind,'ocr_backend':backend if ocr!='off' else 'off','visuals':visuals,'ocr_all_regions':ocr_all_regions,'pages':[]}
    # The runtime owns only a fresh task-local directory, and removes it at exit.
    with tempfile.TemporaryDirectory(prefix='.pdf-layout-',dir=output.parent) as temp:
        work=Path(temp);engine=make_rapid() if ocr!='off' and backend=='rapid' else None
        with normalise(source,work) as pdf:
            if kind=='pptx':
                sizes={(round(p.rect.width,2),round(p.rect.height,2)) for p in pdf}
                if len(sizes)>1:raise ValueError('Mixed PDF page sizes cannot map exactly to one PPT canvas. Normalise sizes first.')
                doc=Presentation();doc.slide_width=Pt(pdf[0].rect.width);doc.slide_height=Pt(pdf[0].rect.height)
            else:
                doc=Document()
                for p in list(doc.paragraphs):p._element.getparent().remove(p._element)
                word.setup_section(doc,pdf[0].rect.width,pdf[0].rect.height,0,0)
            for page in pdf:
                lines,art,record=reconstruct(page,work,engine,ocr,backend,visuals,ocr_all_regions);report['pages'].append(record)
                w,h=record['size']
                if kind=='docx':
                    if page.number and tuple(record['size'])!=tuple(report['pages'][page.number-1]['size']):
                        sec=doc.add_section(WD_SECTION.NEW_PAGE);sec.page_width=WPt(w);sec.page_height=WPt(h)
                        holder=word.add_blank_paragraph(doc)
                    else:holder=word.add_blank_paragraph(doc,page_break=page.number>0)
                    if visuals!='none':
                        pictures.add_full_page_picture(holder,str(art),w,h)
                        anchors=holder._p.findall('.//'+qn('wp:anchor'));anchors[-1].set('behindDoc','1')
                    for line in lines:
                        x,y,bw,bh,lh,align=textbox_geometry(line,w,h)
                        runs=''
                        for s in line['spans']:
                            color=s['color'];font=word.norm_font(s['font']) if line['origin']=='native' else s['font']
                            rgb=RGBColor((color>>16)&255,(color>>8)&255,color&255)
                            if visuals=='none' and color==0xFFFFFF:rgb=RGBColor(31,31,31)
                            runs+=word.run_xml(s['text'],font,round(s['size'],1),rgb,bold=bool(s.get('flags',0)&16) or 'Bold' in s['font'])
                        inner=word.para_xml(runs,align=align,line_pt=lh)
                        drawing=parse_xml(word.textbox_xml(doc.part.next_id,x,y,bw,bh,inner))
                        run=parse_xml('<w:r %s/>'%nsdecls('w'));run.append(drawing);holder._p.append(run)
                else:
                    slide=doc.slides.add_slide(doc.slide_layouts[6])
                    if visuals!='none':slide.shapes.add_picture(str(art),0,0,Pt(w),Pt(h))
                    for line in lines:
                        x,y,bw,bh,lh,align=textbox_geometry(line,w,h)
                        box=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(bw),Pt(bh));tf=box.text_frame
                        tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Pt(0);tf.vertical_anchor=MSO_ANCHOR.TOP
                        p=tf.paragraphs[0];p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
                        p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(lh)
                        for s in line['spans']:
                            r=p.add_run();r.text=s['text'];r.font.size=Pt(round(s['size'],1));r.font.name=ppt.norm_font(s['font']) if line['origin']=='native' else s['font'];ppt.set_ea_font(r,r.font.name)
                            r.font.bold=bool(s.get('flags',0)&16) or 'Bold' in s['font'];r.font.italic=bool(s.get('flags',0)&2)
                            c=s['color'];r.font.color.rgb=PColor((c>>16)&255,(c>>8)&255,c&255)
                print(f'page {page.number+1}/{len(pdf)}: {record["native_lines"]} native + {record["ocr_lines"]} OCR lines',flush=True)
            doc.save(str(output))
    report['conversion_seconds']=round(time.perf_counter()-t,3);report['output_sha256']=digest(output)
    report['review_required']=any(p['ocr_lines'] for p in report['pages'])
    report['editability']='native/OCR text editable; retained pictures, logos, rules and diagrams are raster graphics'
    Path(str(output)+'.conversion.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return report
