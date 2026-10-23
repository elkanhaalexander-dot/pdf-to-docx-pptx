"""Separate page/text checks, visual preservation, and OCR review status."""
from pathlib import Path
import argparse,collections,json,tempfile,zipfile,xml.etree.ElementTree as ET
from layout_common import normalise,render_office,digest,norm,native_lines
import pymupdf as fitz

def raw_pages(path):
    if path.suffix.lower()=='.pptx':
        from pptx import Presentation
        return [''.join(s.text_frame.text for s in sl.shapes if s.has_text_frame) for sl in Presentation(path).slides]
    W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    with zipfile.ZipFile(path) as z:root=ET.fromstring(z.read('word/document.xml'))
    pages=['']
    for block in root.find(W+'body'):
        if block.tag==W+'p':
            if block.find('./'+W+'pPr/'+W+'pageBreakBefore') is not None:pages.append('')
            for child in block:
                if child.tag==W+'pPr':continue
                for el in child.iter():
                    if el.tag==W+'t':pages[-1]+=el.text or ''
                    if el.tag==W+'br' and el.get(W+'type')=='page':pages.append('')
        elif block.tag==W+'tbl':pages[-1]+=''.join(e.text or '' for e in block.iter(W+'t'))
    return pages

def visual_difference(source,target,ocr_regions):
    import numpy as np
    if abs(source.rect.width-target.rect.width)>.2 or abs(source.rect.height-target.rect.height)>.2:return {'error':'page dimensions differ'}
    a=source.get_pixmap(alpha=False);b=target.get_pixmap(alpha=False)
    aa=np.frombuffer(a.samples,dtype=np.uint8).reshape(a.height,a.width,3)
    bb=np.frombuffer(b.samples,dtype=np.uint8).reshape(b.height,b.width,3)
    if aa.shape!=bb.shape:
        if abs(aa.shape[0]-bb.shape[0])>1 or abs(aa.shape[1]-bb.shape[1])>1:return {'error':'raster dimensions differ'}
        from PIL import Image
        bb=np.array(Image.fromarray(bb).resize((a.width,a.height)))
    mask=np.ones(aa.shape[:2],dtype=bool)
    boxes=[s['bbox'] for l in native_lines(source) for s in l['spans']]+ocr_regions
    for x0,y0,x1,y1 in boxes:mask[max(0,int(y0)-4):min(a.height,int(y1)+5),max(0,int(x0)-4):min(a.width,int(x1)+5)]=False
    changed=(np.max(np.abs(aa.astype(float)-bb.astype(float)),axis=2)>35)&mask
    return {'changed_nontext_fraction':round(float(changed.sum()/max(1,mask.sum())),6),'checked_pixels':int(mask.sum())}

def verify(source,product,render=False,max_diff=.02):
    source=Path(source).resolve();product=Path(product).resolve()
    result={'checks':[],'warnings':[],'passed':False,'ocr_accuracy_verified':False}
    meta_path=Path(str(product)+'.conversion.json');meta=None
    if meta_path.exists():
        meta=json.loads(meta_path.read_text(encoding='utf-8'))
        if meta.get('source_sha256')!=digest(source) or meta.get('output_sha256')!=digest(product):
            result['checks'].append({'check':'provenance','passed':False,'detail':'source or product changed since conversion'});return result
    else:result['warnings'].append('No metadata: OCR regions and image coverage are unknown.')
    with tempfile.TemporaryDirectory(prefix='.pdf-verify-',dir=product.parent) as temp:
        work=Path(temp)
        with normalise(source,work) as original:
            expected=[norm(p.get_text()) for p in original];actual=raw_pages(product)
            result['checks'].append({'check':'raw_page_count','passed':len(expected)==len(actual),'source_pages':len(expected),'product_pages':len(actual)})
            missing=[]
            for i,want in enumerate(expected):
                absent=collections.Counter(want)-collections.Counter(norm(actual[i]) if i<len(actual) else '')
                if absent:missing.append({'page':i+1,'missing_chars':sum(absent.values())})
            result['checks'].append({'check':'native_text_per_page','passed':not missing,'missing':missing,'scope':'Native PDF text only; not an OCR accuracy measure.'})
            if meta:
                result['ocr_text_lines']=sum(p['ocr_lines'] for p in meta['pages'])
                result['ocr_review_candidates']=sum(len(p['review']) for p in meta['pages'])
                result['warnings'].append('OCR text is editable but needs comparison against source images. Text-layer checks cannot establish OCR accuracy.')
                result['warnings'].append('Retained graphics and table rules are raster; OCR text boxes are not native Word table cells.')
            if render:
                pdf=render_office(product,work/'rendered',work/'lo-profile');rendered=fitz.open(pdf)
                result['checks'].append({'check':'rendered_page_count','passed':len(rendered)==len(original),'rendered_pages':len(rendered)})
                clipped=[];visual=[];lost=[];lost_ocr=[]
                for i,p in enumerate(rendered):
                    if i>=len(original):break
                    absent=collections.Counter(expected[i])-collections.Counter(norm(p.get_text()))
                    if absent:lost.append({'page':i+1,'missing_chars':sum(absent.values())})
                    if meta:
                        want=norm(meta['pages'][i]['native_text']+meta['pages'][i]['ocr_text'])
                        missing_ocr=collections.Counter(want)-collections.Counter(norm(p.get_text()))
                        if missing_ocr:lost_ocr.append({'page':i+1,'missing_chars':sum(missing_ocr.values())})
                    for b in p.get_text('dict')['blocks']:
                        if b['type']==0:
                            x0,y0,x1,y1=b['bbox']
                            if x0< -2 or y0< -2 or x1>p.rect.width+2 or y1>p.rect.height+2:clipped.append(i+1)
                    regions=meta['pages'][i]['ocr_regions'] if meta and i<len(meta['pages']) else []
                    visual.append({'page':i+1,**visual_difference(original[i],p,regions)})
                result['checks'].append({'check':'rendered_native_text','passed':not lost,'missing':lost})
                if meta:result['checks'].append({'check':'rendered_native_and_ocr_text','passed':not lost_ocr,'missing':lost_ocr,'scope':'Checks OCR text renders; does not establish recognition accuracy.'})
                result['checks'].append({'check':'text_out_of_bounds','passed':not clipped,'pages':sorted(set(clipped))})
                bad=[p for p in visual if p.get('error') or p.get('changed_nontext_fraction',1)>max_diff]
                result['checks'].append({'check':'nontext_visual_preservation','passed':not bad,'threshold':max_diff,'pages':visual,'bad_pages':[p['page'] for p in bad],
                    'scope':'Outside native/OCR text boxes only; does not prove OCR accuracy or identical font rendering.'})
                rendered.close()
            else:result['warnings'].append('NOT RENDERED: actual pagination, clipping and visual preservation were not checked.')
    result['passed']=all(c['passed'] for c in result['checks'])
    result['status']=('MECHANICAL_CHECKS_PASS_OCR_REVIEW_REQUIRED' if result.get('ocr_text_lines',0) else 'CHECKS_PASS') if result['passed'] else 'FAIL'
    if result['passed'] and not render:result['status']='TEXT_STRUCTURE_PASS_RENDER_REQUIRED'
    return result

def main():
    ap=argparse.ArgumentParser();ap.add_argument('source');ap.add_argument('product');ap.add_argument('--render',action='store_true');ap.add_argument('--json',dest='json_path');ap.add_argument('--max-visual-diff',type=float,default=.02)
    args=ap.parse_args();result=verify(args.source,args.product,args.render,args.max_visual_diff)
    text=json.dumps(result,ensure_ascii=False,indent=2)
    if args.json_path:Path(args.json_path).write_text(text,encoding='utf-8')
    print(text);return 0 if result['passed'] else 1
if __name__=='__main__':raise SystemExit(main())
