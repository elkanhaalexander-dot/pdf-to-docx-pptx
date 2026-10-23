# -*- coding: utf-8 -*-
"""Editable Word: preserve native text, OCR remaining text, retain artwork.

The default slide profile uses layout_common for PDF/image/PPT inputs.
The legacy flow profile requires --ocr off and does not preserve artwork/pages.
"""
import argparse
import os
import sys
sys.dont_write_bytecode=True

HERE = os.path.dirname(os.path.abspath(__file__))
DEPS = os.path.join(os.path.dirname(HERE), 'vendor')
if DEPS not in sys.path:
    sys.path.insert(0, DEPS)

import pymupdf  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING  # noqa: E402
from docx.oxml.ns import qn, nsdecls  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402

EMU_PER_PT = 12700
WHITE_FALLBACK = RGBColor(0x1F, 0x1F, 0x1F)

# python-docx 的 nsmap 里没有 wps（WordprocessingShape），手工补齐
NS_WPS = ('xmlns:wps="http://schemas.microsoft.com/office/word/2010/'
          'wordprocessingShape"')

# PDF 里的字体名 → Word 可识别的字体名
FONT_MAP = {
    'SimSun': 'SimSun',
    'TimesNewRomanPS-BoldMT': 'Times New Roman',
    'TimesNewRomanPSMT': 'Times New Roman',
    'Times New Roman': 'Times New Roman',
    'Arial': 'Arial',
    'Arial-BoldMT': 'Arial',
}


def norm_font(name):
    base = name.split(',')[0].strip()
    if base in FONT_MAP:
        return FONT_MAP[base]
    low = base.lower()
    if 'simsun' in low or 'song' in low:
        return 'SimSun'
    if 'times' in low:
        return 'Times New Roman'
    if 'arial' in low:
        return 'Arial'
    if 'calibri' in low:
        return 'Calibri'
    # 等线 / DengXian 在 PDF 里是乱码名，统一映射到「等线」
    return 'DengXian'


def span_props(s):
    """返回 (字体, 字号pt, RGBColor)。白色文字替换为深灰，因为背景已去除。"""
    size = round(s['size'], 1)
    color = s['color']
    if color == 0xFFFFFF:
        rgb = WHITE_FALLBACK
    else:
        rgb = RGBColor((color >> 16) & 0xFF, (color >> 8) & 0xFF, color & 0xFF)
    return norm_font(s['font']), size, rgb


def style_run(run, font, size, rgb, bold=False):
    run.font.size = Pt(size)
    run.font.color.rgb = rgb
    run.font.bold = bold
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml('<w:rFonts %s/>' % nsdecls('w'))
        rPr.insert(0, rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rFonts.set(qn(attr), font)


# ---------------------------------------------------------------- 文本框 XML
def textbox_xml(pid, x, y, w, h, inner_paragraphs_xml, anchor='t'):
    return (
        '<w:drawing %s>'
        '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        'relativeHeight="251659264" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="page"><wp:posOffset>%d</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="page"><wp:posOffset>%d</wp:posOffset></wp:positionV>'
        '<wp:extent cx="%d" cy="%d"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:wrapNone/>'
        '<wp:docPr id="%d" name="Text %d"/>'
        '<wp:cNvGraphicFramePr/>'
        '<a:graphic %s>'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp %s %s>'
        '<wps:cNvSpPr txBox="1"/>'
        '<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:noFill/><a:ln><a:noFill/></a:ln></wps:spPr>'
        '<wps:txbx><w:txbxContent>%s</w:txbxContent></wps:txbx>'
        '<wps:bodyPr rot="0" vert="horz" wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" '
        'anchor="%s" anchorCtr="0"><a:noAutofit/></wps:bodyPr>'
        '</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>'
    ) % (nsdecls('w', 'wp', 'a', 'r'),
         int(round(x * EMU_PER_PT)), int(round(y * EMU_PER_PT)),
         int(round(w * EMU_PER_PT)), int(round(h * EMU_PER_PT)),
         pid, pid, nsdecls('a'), nsdecls('w'), NS_WPS,
         int(round(w * EMU_PER_PT)), int(round(h * EMU_PER_PT)),
         inner_paragraphs_xml, anchor)


def para_xml(runs_xml, align='left', line_pt=1.0, space_before=0.0, indent_pt=0.0):
    jc = {'left': 'left', 'center': 'center', 'right': 'right'}[align]
    ind = ''
    if indent_pt > 0.5:
        ind = '<w:ind w:left="0" w:firstLine="%d"/>' % int(round(indent_pt * 20))
    return (
        '<w:p %s><w:pPr>'
        '<w:spacing w:before="%d" w:after="0" w:line="%d" w:lineRule="exact"/>%s'
        '<w:jc w:val="%s"/>'
        '<w:rPr><w:sz w:val="2"/></w:rPr>'
        '</w:pPr>%s</w:p>'
    ) % (nsdecls('w'), int(round(space_before * 20)), int(round(line_pt * 20)),
         ind, jc, runs_xml)


def run_xml(text, font, size, rgb, bold=False):
    esc = (text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'))
    hexcolor = '%02X%02X%02X' % (rgb[0], rgb[1], rgb[2])
    b = '<w:b/>' if bold else ''
    return (
        '<w:r %s><w:rPr>'
        '<w:rFonts w:ascii="%s" w:hAnsi="%s" w:eastAsia="%s" w:cs="%s"/>%s'
        '<w:color w:val="%s"/><w:sz w:val="%d"/><w:szCs w:val="%d"/>'
        '</w:rPr><w:t xml:space="preserve">%s</w:t></w:r>'
    ) % (nsdecls('w'), font, font, font, font, b, hexcolor,
         int(round(size * 2)), int(round(size * 2)), esc)


def add_blank_paragraph(doc, page_break=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(1)
    if page_break:
        pf.page_break_before = True
    return p


def setup_section(doc, page_w, page_h, margin_l=0, margin_t=0):
    sec = doc.sections[0]
    sec.page_width = Pt(page_w)
    sec.page_height = Pt(page_h)
    sec.left_margin = Pt(margin_l)
    sec.right_margin = Pt(max(page_w - margin_l - (page_w - 2 * margin_l), 0))
    sec.top_margin = Pt(margin_t)
    sec.bottom_margin = Pt(0)
    sec.header_distance = Pt(0)
    sec.footer_distance = Pt(0)
    sectPr = sec._sectPr
    for tag in ('w:headerReference', 'w:footerReference'):
        for el in sectPr.findall(qn(tag)):
            sectPr.remove(el)
    normal = doc.styles['Normal']
    normal.font.size = Pt(1)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(1)
    return sec


# ---------------------------------------------------------------- slide 模式
def build_slide(pdf_path, out_path):
    """每「视觉行」生成一个独立的浮动文本框，按原 PDF 的坐标定位。

    为什么不按 text-block 建框：PPT 导出的 PDF 里，同一个 block 常包含横向并列的
    若干行（例如「稳」「词性」「搭配」「要得当」），它们 x 起点不同、并非上下关系。
    若塞进一个框按段落堆叠，就会被挤到框外裁掉。逐行建框可让每段文字的位置
    与原 PDF 完全一致，且每段都是可独立编辑的文本。
    """
    src = pymupdf.open(pdf_path)
    doc = Document()
    body = doc.element.body
    for p in list(doc.paragraphs):
        body.remove(p._element)
    page_w, page_h = src[0].rect.width, src[0].rect.height
    setup_section(doc, page_w, page_h, 0, 0)

    pid = 1000
    for pno in range(src.page_count):
        page = src[pno]
        holder = add_blank_paragraph(doc, page_break=(pno > 0))
        for block in page.get_text('dict')['blocks']:
            if block['type'] != 0:
                continue
            for ln in block['lines']:
                spans = [s for s in ln['spans'] if s['text'].strip()]
                if not spans:
                    continue
                text = ''.join(s['text'] for s in ln['spans'])
                if not text.strip():
                    continue

                lx0, ly0, lx1, ly1 = ln['bbox']
                lw = lx1 - lx0
                size = max(round(s['size'], 1) for s in spans)
                ln_h = max(ly1 - ly0, size * 1.02)

                # 对齐方式按「行中心 vs 页面中心」判断
                center = (lx0 + lx1) / 2.0
                if abs(center - page_w / 2.0) < 12 and lw < page_w * 0.9:
                    align = 'center'
                elif lx1 > page_w - 60 and lx0 > page_w * 0.5:
                    align = 'right'
                else:
                    align = 'left'

                # 宽度留余量（Word 字宽与 PDF 度量不同），但锚点按对齐方式固定
                want_w = lw * 1.4 + 30
                if align == 'center':
                    w = min(want_w, page_w)
                    x = center - w / 2.0
                elif align == 'right':
                    w = min(want_w, page_w)
                    x = (lx1 + 8) - w
                else:
                    x = lx0
                    w = min(want_w, page_w - x)
                x = max(0.0, min(x, max(page_w - w, 0)))
                w = max(min(w, page_w - x), 1.0)
                y = max(ly0 - 2, 0)
                h = min(ln_h + 8, page_h - y)

                runs = ''.join(run_xml(s['text'], *span_props(s)) for s in spans)
                inner = para_xml(runs, align=align, line_pt=ln_h)
                pid += 1
                drawing = parse_xml(textbox_xml(pid, x, y, w, h, inner))
                run_el = parse_xml('<w:r %s/>' % nsdecls('w'))
                run_el.append(drawing)
                holder._p.append(run_el)

    doc.save(out_path)
    src.close()
    return out_path


# ---------------------------------------------------------------- flow 模式
def group_rows(page, tol=4.0):
    """把页面上的所有文字行按纵向位置聚成「视觉行」，行内按 x 排序。"""
    items = []
    for block in page.get_text('dict')['blocks']:
        if block['type'] != 0:
            continue
        for ln in block['lines']:
            txt = ''.join(s['text'] for s in ln['spans'])
            if not txt.strip():
                continue
            items.append({'y': ln['bbox'][1], 'x': ln['bbox'][0], 'line': ln})
    items.sort(key=lambda it: (it['y'], it['x']))
    rows = []
    for it in items:
        if rows and abs(it['y'] - rows[-1]['y']) <= tol:
            rows[-1]['items'].append(it)
        else:
            rows.append({'y': it['y'], 'items': [it]})
    for r in rows:
        r['items'].sort(key=lambda it: it['x'])
    return rows


def build_flow(pdf_path, out_path, margin_l=72, margin_t=72):
    src = pymupdf.open(pdf_path)
    doc = Document()
    body = doc.element.body
    for p in list(doc.paragraphs):
        body.remove(p._element)
    page_w, page_h = src[0].rect.width, src[0].rect.height
    setup_section(doc, page_w, page_h, margin_l, margin_t)

    for pno in range(src.page_count):
        page = src[pno]
        rows = group_rows(page)
        prev_bottom = None
        for row in rows:
            items = row['items']
            # 行高与字号：取该视觉行最大字号
            size = max(round(s['size'], 1) for it in items
                       for s in it['line']['spans'] if s['text'].strip())
            # 段前间距：与上一视觉行的实际间隙
            gap = 0.0
            if prev_bottom is not None:
                gap = max(row['y'] - prev_bottom, 0.0)
            prev_bottom = max(it['line']['bbox'][3] for it in items)

            p = add_blank_paragraph(doc, page_break=(pno > 0 and row is rows[0]))
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(size * 1.32)
            pf.space_before = Pt(min(gap, size * 2.0))

            # 居中的标题行
            x0 = min(it['x'] for it in items)
            x1 = max(it['line']['bbox'][2] for it in items)
            if abs((x0 + x1) / 2.0 - page_w / 2.0) < 15 and x0 > margin_l + 30:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 制表位：把各列起点换算成相对左缩进的偏移
            stops = sorted({it['x'] for it in items})
            for sx in stops[1:]:
                offset = sx - margin_l
                if offset > 0:
                    pf.tab_stops.add_tab_stop(Pt(offset))

            first = True
            for it in items:
                if not first:
                    p.add_run('\t')
                first = False
                for s in it['line']['spans']:
                    if not s['text']:
                        continue
                    font, sz, rgb = span_props(s)
                    bold = 'Bold' in s['font']
                    style_run(p.add_run(s['text']), font, sz, rgb, bold)

    doc.save(out_path)
    src.close()
    return out_path


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument('pdf')
    ap.add_argument('docx')
    ap.add_argument('--profile', choices=['slide', 'flow'], default='slide')
    ap.add_argument('--margin-l', type=float, default=72)
    ap.add_argument('--margin-t', type=float, default=72)
    ap.add_argument('--ocr', choices=['auto','off'], default='auto', help='OCR residual image/outlined text; native PDF text is always retained')
    ap.add_argument('--ocr-backend', choices=['rapid','windows'], default='rapid')
    ap.add_argument('--ocr-all-regions',action='store_true',help='Also OCR small corner logos/footer labels; use if these corners contain editable body text')
    ap.add_argument('--visuals', choices=['preserve','none'], default='preserve', help='preserve artwork by default; none explicitly requests text-only output')
    args = ap.parse_args(argv[1:])
    if args.profile == 'slide':
        from layout_common import convert
        convert(args.pdf, args.docx, 'docx', args.ocr, args.ocr_backend, args.visuals,args.ocr_all_regions)
    else:
        if args.ocr!='off':
            ap.error('flow is a legacy native-text-only reflow mode. Use --ocr off, or slide for editable OCR and page correspondence.')
        build_flow(args.pdf, args.docx, args.margin_l, args.margin_t)
        print('WARNING: flow does not preserve images/outlined text or exact page correspondence.')
    print('OK -> %s' % args.docx)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
